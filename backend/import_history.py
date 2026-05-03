"""Import every transaction-bearing row from the master sheet into SQLite.

Sources:
- GRN tab                         → 'grn'
- Sheet26 (older GRN-style log)   → 'sheet26' (deduped vs grn)
- Daily blocks in stock sheets    → 'daily_recv' / 'daily_issue'
- Cans Labels per-row issues      → 'cans_label'
- Thread & Tags per-row events    → 'thread_tag'

All non-dashboard rows are audit-only (do not mutate Item.current_stock).
Idempotent via dedup_key.
"""
import hashlib
import os
import re
from datetime import datetime
from openpyxl import load_workbook

from .database import Base, engine, SessionLocal
from . import models

SRC = os.path.join(os.path.dirname(__file__), "..", "data", "cache", "master.xlsx")


def parse_date(v):
    if isinstance(v, datetime):
        return v
    if not v:
        return None
    s = str(v).strip()
    for fmt in ("%d-%m-%Y", "%d/%m/%Y", "%Y-%m-%d", "%d-%m-%y", "%d/%m/%y"):
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            pass
    return None


def num(v):
    try:
        if v is None or v == "":
            return None
        return float(v)
    except (ValueError, TypeError):
        return None


def s(v):
    return str(v).strip() if v is not None and str(v).strip() else None


def norm(v):
    return re.sub(r"\s+", " ", str(v or "").strip().lower())


def dedup_key(*parts) -> str:
    return hashlib.md5("|".join(norm(p) for p in parts).encode()).hexdigest()


def get_or_create_item(db, name, category, unit="PCS"):
    item = db.query(models.Item).filter(models.Item.name.ilike(name)).first()
    if item:
        return item
    nm = norm(name)
    for it in db.query(models.Item).filter(models.Item.category == category).all():
        if norm(it.name) == nm:
            return it
    item = models.Item(
        category=category or "consumable",
        name=name,
        unit=unit or "PCS",
        notes="auto-created from history import",
    )
    db.add(item)
    db.flush()
    return item


CATEGORY_MAP = {
    "rm": "raw_material", "pm": "packaging", "cm": "packaging",
    "labels": "label", "label": "label", "stickers": "label",
    "fa": "finished_good", "liquid": "consumable",
    "pantry": "housekeeping", "stationary": "housekeeping", "stationery": "housekeeping",
    "maintenance": "housekeeping", "hk/stationary": "housekeeping",
    "h+s+p": "housekeeping", "housekeeping +stationary+pantry": "housekeeping",
    "thread": "thread_tag", "tag": "thread_tag",
    "rubber band": "consumable", "coin": "consumable", "electric": "housekeeping",
    "dn": "consumable",
}


def map_category(raw):
    if not raw:
        return "consumable"
    return CATEGORY_MAP.get(str(raw).strip().lower(), "consumable")


# ---------- GRN ----------
def import_grn(db, ws):
    n = 0
    for row_idx, r in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        invoice = s(r[0])
        desc = s(r[2])
        if not desc:
            continue
        category = map_category(s(r[3]))
        unit = s(r[4]) or "PCS"
        bill_qty = num(r[5])
        recv_qty = num(r[6])
        short = num(r[7])
        recv_date = parse_date(r[8]) or parse_date(r[1])
        vendor = s(r[9])
        person = s(r[10])
        status = s(r[12])
        remarks = s(r[13])
        qty = recv_qty if recv_qty is not None else bill_qty
        if qty is None or qty <= 0:
            continue
        item = get_or_create_item(db, desc, category, unit)
        key = dedup_key("grn", row_idx, invoice, desc, recv_date, qty, vendor)
        if db.query(models.Transaction).filter_by(dedup_key=key).first():
            continue
        db.add(models.Transaction(
            item_id=item.id, txn_type="receive", qty=qty,
            txn_date=recv_date or datetime(2025, 1, 1),
            party=vendor, person=person, invoice_no=invoice,
            remarks=" | ".join(x for x in [status, remarks] if x) or None,
            bill_qty=bill_qty, short_excess=short,
            source="grn", dedup_key=key,
        ))
        n += 1
    return n


# ---------- Sheet26 ----------
def import_sheet26(db, ws):
    n = 0
    for row_idx, r in enumerate(ws.iter_rows(min_row=3, values_only=True), start=3):
        po = s(r[0])
        desc = s(r[2])
        if not desc:
            continue
        category = map_category(s(r[3]))
        unit = s(r[4]) or "PCS"
        bill_qty = num(r[5])
        recv_qty = num(r[6])
        short = num(r[7])
        recv_date = parse_date(r[8])
        vendor = s(r[9])
        person = s(r[10])
        qty = recv_qty if recv_qty is not None else bill_qty
        if qty is None or qty <= 0:
            continue
        item = get_or_create_item(db, desc, category, unit)
        # Soft dedup vs GRN: same invoice, item, qty
        existing = db.query(models.Transaction).filter(
            models.Transaction.invoice_no == po,
            models.Transaction.item_id == item.id,
            models.Transaction.qty == qty,
            models.Transaction.source == "grn",
        ).first()
        if existing:
            continue
        key = dedup_key("sheet26", row_idx, po, desc, recv_date, qty, vendor)
        if db.query(models.Transaction).filter_by(dedup_key=key).first():
            continue
        db.add(models.Transaction(
            item_id=item.id, txn_type="receive", qty=qty,
            txn_date=recv_date or datetime(2025, 1, 1),
            party=vendor, person=person, invoice_no=po,
            bill_qty=bill_qty, short_excess=short,
            source="sheet26", dedup_key=key,
        ))
        n += 1
    return n


# ---------- Daily blocks ----------
DAILY_LAYOUTS = {
    "RM": dict(item_col=1, unit_col=3, first_day=5, block=7,
               offsets={"recv": 1, "recv_by": 2, "issue": 3, "iss_by": 4, "iss_to": 5},
               default_unit="KG", category="raw_material"),
    "PACKAGING": dict(item_col=2, unit_col=3, first_day=5, block=7,
                      offsets={"recv": 1, "recv_by": 2, "issue": 3, "iss_by": 4, "iss_to": 5},
                      default_unit="PCS", category="packaging"),
    "CONSUMABLE": dict(item_col=2, unit_col=3, first_day=5, block=7,
                       offsets={"recv": 1, "recv_by": 2, "issue": 3, "iss_by": 4, "iss_to": 5},
                       default_unit="PCS", category="consumable"),
    "Labels": dict(item_col=1, unit_col=2, first_day=4, block=7,
                   offsets={"recv": 1, "recv_by": 2, "issue": 3, "iss_by": 4, "iss_to": 5},
                   default_unit="PCS", category="label"),
    "Housekeeping +Stationary+Pantry": dict(
        item_col=3, unit_col=5, first_day=6, block=5,
        offsets={"recv": 1, "issue": 2, "iss_by": 3},
        default_unit="PCS", category="housekeeping"),
}


def import_daily_blocks(db, ws, sheet_label, header_row=3, date_row=2):
    cfg = DAILY_LAYOUTS[sheet_label]
    date_row_vals = next(ws.iter_rows(min_row=date_row, max_row=date_row, values_only=True))
    date_cols = [(i, v) for i, v in enumerate(date_row_vals) if isinstance(v, datetime)]
    n_recv, n_iss = 0, 0
    for row_idx, r in enumerate(ws.iter_rows(min_row=header_row + 1, values_only=True),
                                start=header_row + 1):
        item_name = s(r[cfg["item_col"]]) if cfg["item_col"] < len(r) else None
        if not item_name:
            continue
        unit = (s(r[cfg["unit_col"]]) if cfg["unit_col"] < len(r) else None) or cfg["default_unit"]
        item = get_or_create_item(db, item_name, cfg["category"], unit)
        for i, day in date_cols:
            off = cfg["offsets"]
            recv = r[i + off["recv"]] if (i + off["recv"]) < len(r) else None
            if isinstance(recv, (int, float)) and recv > 0:
                key = dedup_key("d_recv", sheet_label, row_idx, i, day, recv)
                if not db.query(models.Transaction).filter_by(dedup_key=key).first():
                    person = s(r[i + off["recv_by"]]) if "recv_by" in off and (i + off["recv_by"]) < len(r) else None
                    db.add(models.Transaction(
                        item_id=item.id, txn_type="receive", qty=float(recv),
                        txn_date=day, person=person,
                        source="daily_recv", dedup_key=key,
                    ))
                    n_recv += 1
            iss = r[i + off["issue"]] if (i + off["issue"]) < len(r) else None
            if isinstance(iss, (int, float)) and iss > 0:
                key = dedup_key("d_iss", sheet_label, row_idx, i, day, iss)
                if not db.query(models.Transaction).filter_by(dedup_key=key).first():
                    iss_by = s(r[i + off["iss_by"]]) if "iss_by" in off and (i + off["iss_by"]) < len(r) else None
                    iss_to = s(r[i + off["iss_to"]]) if "iss_to" in off and (i + off["iss_to"]) < len(r) else None
                    db.add(models.Transaction(
                        item_id=item.id, txn_type="issue", qty=float(iss),
                        txn_date=day, party=iss_to, person=iss_by,
                        source="daily_issue", dedup_key=key,
                    ))
                    n_iss += 1
    return n_recv, n_iss


# ---------- Cans Labels (per-row issue blocks of 4 cols) ----------
def import_cans_labels(db, ws):
    n = 0
    for row_idx, r in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        name = s(r[2])
        if not name:
            continue
        item = get_or_create_item(db, name, "can_label", "Pcs")
        # Issue blocks at col 7, 11, 15, 19, ... (each: date, qty, vendor, closing)
        col = 7
        while col + 3 < len(r):
            dt = parse_date(r[col])
            qty = num(r[col + 1])
            vendor = s(r[col + 2])
            if dt and qty and qty > 0:
                key = dedup_key("cans", row_idx, col, dt, qty, vendor)
                if not db.query(models.Transaction).filter_by(dedup_key=key).first():
                    db.add(models.Transaction(
                        item_id=item.id, txn_type="issue", qty=qty,
                        txn_date=dt, party=vendor,
                        source="cans_label", dedup_key=key,
                    ))
                    n += 1
            col += 4
    return n


# ---------- Thread & Tags ----------
def import_thread_tags(db, ws):
    """Two parallel logs side by side. Layout per row from R5 onward:
    Left  cols 0..6:  ItemName | RecvKg | RecvPcs | IssueKg | IssuePcs | IssueDate | OrderNo
    Right cols 8..13: ItemName | RecvQty | RecvDate | IssueQty | IssueDate | OrderNo
    """
    n = 0
    for row_idx, r in enumerate(ws.iter_rows(min_row=5, values_only=True), start=5):
        # LEFT
        name_l = s(r[0])
        if name_l:
            item = get_or_create_item(db, name_l, "thread_tag", "Pcs")
            recv_pcs = num(r[2])
            iss_pcs = num(r[4])
            iss_date = parse_date(r[5])
            order = s(r[6])
            if recv_pcs and recv_pcs > 0:
                key = dedup_key("tt_l_r", row_idx, recv_pcs, name_l)
                if not db.query(models.Transaction).filter_by(dedup_key=key).first():
                    db.add(models.Transaction(
                        item_id=item.id, txn_type="receive", qty=recv_pcs,
                        txn_date=iss_date or datetime(2025, 1, 1),
                        order_no=order, source="thread_tag", dedup_key=key,
                    ))
                    n += 1
            if iss_pcs and iss_pcs > 0 and iss_date:
                key = dedup_key("tt_l_i", row_idx, iss_pcs, iss_date, order)
                if not db.query(models.Transaction).filter_by(dedup_key=key).first():
                    db.add(models.Transaction(
                        item_id=item.id, txn_type="issue", qty=iss_pcs,
                        txn_date=iss_date, order_no=order,
                        source="thread_tag", dedup_key=key,
                    ))
                    n += 1
        # RIGHT
        name_r = s(r[8]) if len(r) > 8 else None
        if name_r:
            item = get_or_create_item(db, name_r, "thread_tag", "Pcs")
            recv_q = num(r[9]) if len(r) > 9 else None
            recv_d = parse_date(r[10]) if len(r) > 10 else None
            iss_q = num(r[11]) if len(r) > 11 else None
            iss_d = parse_date(r[12]) if len(r) > 12 else None
            order = s(r[13]) if len(r) > 13 else None
            if recv_q and recv_q > 0:
                key = dedup_key("tt_r_r", row_idx, recv_q, recv_d, name_r)
                if not db.query(models.Transaction).filter_by(dedup_key=key).first():
                    db.add(models.Transaction(
                        item_id=item.id, txn_type="receive", qty=recv_q,
                        txn_date=recv_d or datetime(2025, 1, 1),
                        order_no=order, source="thread_tag", dedup_key=key,
                    ))
                    n += 1
            if iss_q and iss_q > 0 and iss_d:
                key = dedup_key("tt_r_i", row_idx, iss_q, iss_d, order)
                if not db.query(models.Transaction).filter_by(dedup_key=key).first():
                    db.add(models.Transaction(
                        item_id=item.id, txn_type="issue", qty=iss_q,
                        txn_date=iss_d, order_no=order,
                        source="thread_tag", dedup_key=key,
                    ))
                    n += 1
    return n


def run():
    Base.metadata.create_all(engine)
    wb = load_workbook(SRC, data_only=True)
    db = SessionLocal()
    counts = {}
    try:
        counts["grn"] = import_grn(db, wb["GRN"])
        db.commit()
        counts["sheet26"] = import_sheet26(db, wb["Sheet26"])
        db.commit()

        for sheet in ["RM", "PACKAGING", "CONSUMABLE", "Labels", "Housekeeping +Stationary+Pantry"]:
            recv, iss = import_daily_blocks(db, wb[sheet], sheet)
            counts[f"{sheet}_recv"] = recv
            counts[f"{sheet}_iss"] = iss
            db.commit()

        counts["cans_labels"] = import_cans_labels(db, wb["Cans Labels"])
        db.commit()
        counts["thread_tags"] = import_thread_tags(db, wb["Thread & Tags"])
        db.commit()

        total = sum(counts.values())
        print(f"Imported {total} historical transactions:")
        for k, v in sorted(counts.items()):
            print(f"  {k}: {v}")
    finally:
        db.close()


if __name__ == "__main__":
    run()
